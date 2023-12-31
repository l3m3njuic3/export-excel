import pandas as pd 
from openpyxl import load_workbook
import time
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import os
import traceback
from PIL import Image
import docx
from docx.shared import Inches
import sys
import requests
from bs4 import BeautifulSoup
import io
import base64
import argparse
            
def retrieve_excel():
    print('[INFO] IT IS NORMAL FOR RETRIEVAL TO FAIL, PLEASE BE PATIENT.')
    print('[INFO] Process will keep trying until successfully retrieved.')
    while True:
        try:
            # delete previous version of the file so that the script doesn't detect the wrong version
            file = './FYP Student Feedback Survey Form - Live Response .xlsx'
            if os.path.exists(file):
                time.sleep(1)
                print('[INFO] Deleting previous versions of the data file...')
                os.remove(file)
            
            url = 'https://entuedu-my.sharepoint.com/:x:/g/personal/nliaw001_e_ntu_edu_sg/EfIPZdtgjtFDiefeP-P2xhgBND-wv1DJdIVzv0gGLNdMNw?e=fx4y94'
            
            time.sleep(2)
            print('[INFO] Starting browser...')
            time.sleep(1)
            print('[INFO] Attempting to retrieve response from server...')
            
            chrome_options = webdriver.ChromeOptions()
            prefs = {'download.default_directory' : os.getcwd()}
            chrome_options.add_experimental_option('prefs', prefs)
            chrome_options.add_experimental_option('excludeSwitches', ['enable-logging'])
            chrome_options.add_argument('--headless') # set headless so that it runs in the background
            
            driver = webdriver.Chrome(options=chrome_options)
            driver.get(url)
            
            # let browser load
            time.sleep(2)
            
            # switch to iframe
            frame = driver.find_element(By.ID, 'WebApplicationFrame')
            driver.switch_to.frame(frame)
            
            # static id
            file_button = WebDriverWait(driver, 8).until(
                EC.presence_of_element_located((By.ID, "id__3"))
            )
            file_button.click()
            print('[INFO] 🟢 First element located successfully. (1/3)')
            
            # dynamic id therefore use absolute XPATH, unstable behavior and not maintainable in the future
            save_as_button = WebDriverWait(driver, 8).until(
                EC.presence_of_element_located((By.XPATH, '/html/body/div[15]/div[5]/div/div/div/div/div/div[2]/div[2]/div/div[2]/div/div[1]/div[1]/button[3]'))
            )
            save_as_button.click()
            print('[INFO] 🟢 Second element located successfully.')
            
            # static id
            download_button = WebDriverWait(driver, 8).until(
                EC.presence_of_element_located((By.ID, 'DownloadACopy'))
            )
            download_button.click()
            print('[INFO] 🟢 Third element located successfully. (2/3)')
            
            # let browser load
            time.sleep(2)
            
            if os.path.exists(file):
                print('[INFO] 🟢 Successfully retrieved response from server. (3/3)')
            else:
                raise Exception()
            
            driver.quit()
            break
        except Exception as e:
            print('[ERROR] 🔴 Failed to retrieve response from server. Trying again...')
            # print(traceback.format_exc())
            
            driver.quit()

def load_excel():
    print("[INFO] Preparing excel sheets...")
    
    filename = 'FYP Student Feedback Survey Form - Live Response .xlsx'

    questions = ['Cohort',
                'Apply suitable load management techniques to achieve a reliable and sustainable power delivery system.', 
                'Able to construct a fair and effective electricity tariff structure.', 
                'Analyse the various causes of poor power quality in distribution systems due to switching events.',
                'Able to apply appropriate mitigation techniques to improve power quality in distribution systems (voltage sags/swells).',
                'Able to apply mitigation techniques to reduce the harmonic distortions in distribution networks.',
                'Able to predict the position of the sun for the purpose of designing renewable solar energy systems.',
                'Able to estimate the amount of solar insolation and understand how the solar panel can be best aligned for maximum energy collection.',
                'Able to compute the current-voltage characteristics of solar cells, modules and arrays and apply mitigation techniques to shading problems.',
                'Able to predict the power and energy of renewable wind resource at various wind speeds, temperatures and altitudes.',
                'Able to understand the basic design of wind turbines and to predict the generated wind power injected into the grid.',
                'Appreciate the need for lifelong learning and self-education in the design and operation of distribution systems with renewable energy generations.',
                'Comments on the strengths of this course (if any):',
                'Comments on the weaknesses of this course (if any):',
                'Other comments on this course (if any):']

    # read the excel file
    # filter out only the needed columns
    df = pd.read_excel(filename, sheet_name="Form1", usecols=questions)
    
    return df   

def write_to_excel(df):
    filename = "21S2AccrSurvey_EEXXXX-Results (ready).xlsx"
    
    # hashmap for constant lookup time
    column = {
        "Strongly Agree": "B",
        "Agree": "D",
        "Neutral": "F",
        "Disagree": "H",
        "Strongly Disagree": "J",
        "Strongly Agree\xa0": "J"
    }

    row = {
        "Question 1": "9",
        "Question 2": "10",
        "Question 3": "11",
        "Question 4": "12",
        "Question 5": "13",
        "Question 6": "14",
        "Question 7": "15",
        "Question 8": "16",
        "Question 9": "17",
        "Question 10": "18",
        "Question 11": "19"
    }

    # load the writer
    workbook = load_workbook(filename)
    worksheets = [workbook['EEXXXX-FTPT'], workbook['EEXXXX-FT'], workbook["EEXXXX-PT"]]
    
    # set all cell values to 0 from all worksheets
    for worksheet in worksheets:
        for c in column.values():
            for r in row.values():
                worksheet[f"{c}{r}"] = 0
            
            # print(f"Cell {c}{r}: Resetted to 0")
    time.sleep(1)
    print("[INFO] Excel file ready. Writing:")
    time.sleep(1)
            
    # write to each cell
    for index, each_row in df.iterrows():
        skipped = False
        for question_no, data in enumerate(each_row, start=0):
            # update comments 
            if data not in column.keys():
                # worksheet to write to
                if data == 'EEE Full-Time':
                    worksheet = worksheets[1]
                elif data == 'EEE Part-Time':
                    worksheet = worksheets[2]
                elif data == 'IEM Full-Time':
                    # skip row if cohort is invalid
                    skipped = True
                    continue
                else:
                    # comments on the course
                    if question_no == 12:
                        column_comments = "A"
                        row_comments = 23
                    elif question_no == 13:
                        column_comments = "A"
                        row_comments = 35
                    elif question_no == 14:
                        column_comments = "A"
                        row_comments = 56
                
                    while True:
                        if worksheet[f"{column_comments}{row_comments}"].value:
                            row_comments += 1
                        else:
                            worksheet[f"{column_comments}{row_comments}"] = data
                            break
            else:
                value = worksheet[f"{column[data]}{row[f'Question {question_no}']}"].value
                worksheet[f"{column[data]}{row[f'Question {question_no}']}"] = int(value) + 1
                    
                # print(f"Cell {column[data]}{row[f'Question {question_no}']}: Updated to {int(value) + 1}")
        if skipped:
            print(f"         Row {index + 1} is invalid. Row skipped.")
        else:
            print(f"         Successfully updated row {index + 1} in {worksheet.title}.")
        
    print("[INFO] Done!")

    try:
        workbook.save(filename)
        print("[INFO] 🟢 Successfully saved file.")
    except PermissionError:
        print("[ERROR] 🔴 Failed to save changes to file.")
        print("[ERROR] 🔴 Please make sure the file is closed before trying again.")

def convert_excel_to_html():
    print('[INFO] Preparing Microsoft Word Document for exporting...')
        
    # Set the URL to which you want to send the file
    url = "https://api.products.aspose.app/cells/conversion/api/ConversionApi/Convert?outputType=HTML"

    # Set the path to the file you want to send
    file_path = os.path.abspath('21S2AccrSurvey_EEXXXX-Results (ready).xlsx')

    # Create a dictionary with the file key and the file object
    files = {"file": open(file_path, "rb")}

    # Send the file using requests.post()
    response = requests.post(url, files=files)
    
    download_file_link = f'https://api.products.aspose.app/cells/conversion/api/Download/{response.json()["FolderName"]}?file=21S2AccrSurvey_EEXXXX-Results%20(ready).html'
    
    file = requests.get(download_file_link, allow_redirects=True)
    
    with open('21S2AccrSurvey_EEXXXX-Results (ready).html', 'wb') as infile:
        infile.write(file.content)
    
    # with open('21S2AccrSurvey_EEXXXX-Results (ready).html', 'r') as infile:
    #     formatted_text = BeautifulSoup(infile, 'html.parser').prettify()
    
    # with open('21S2AccrSurvey_EEXXXX-Results (ready).html', 'w') as outfile:
    #     outfile.write(formatted_text)
    
    print('[INFO] Done!')

def extract_chart_to_docx():
    print("[INFO] Extracting charts...")
    filename = "21S2AccrSurvey_EEXXXX-Results (ready).html"
    
    with open(filename, 'r') as infile:
        file_content = infile.read()
    
    soup = BeautifulSoup(file_content, 'html.parser')
    
    images = soup.find_all('img')

    image_names = ['image_0.png', 'image_1.png', 'image_2.png',]
    
    for index, image in enumerate(images):
        base64_string = image['src'][22:]
        
        # Remove the header from the base64 string
        image_data = base64.b64decode(base64_string)
        
        # Open the image data as a PIL Image object
        image = Image.open(io.BytesIO(image_data))
        
        # Save the image as a PNG file
        image.save(f'{image_names[index]}', 'PNG')
    
    doc = docx.Document()
    
    for image_name in image_names:
        doc.add_picture(image_name, width=Inches(6), height=Inches(4))
    
    doc.save('output.docx')
    
    print("[INFO] 🟢 Output saved to './output.docx'.")
    
    clean_folder(image_names, filename)

def clean_folder(image_names, filename):
    for image_name in image_names:
        os.remove(image_name)
    os.remove(filename)
    
def main(skip):
    if not skip:
        retrieve_excel()
    
    df = load_excel()
    write_to_excel(df)
    convert_excel_to_html()
    extract_chart_to_docx()

def test():
    retrieve_excel()
    # df = load_excel()
    # write_to_excel(df)
    # extract_chart_to_docx()
    # convert_excel_to_html()

if __name__ == "__main__": 
    parser = argparse.ArgumentParser()
    parser.add_argument('-s', '--skip', action='store_true', help='Skip retrieval of responses.')
    args = parser.parse_args()
    skip = False
    
    if args.skip:
        skip = True
        main(skip)
    
    if len(sys.argv) > 1:
        if sys.argv[1] == 'dev':
            print('Test Dev Block')
            test()
    else:
        main(skip)

    
